#!/usr/bin/env python3
"""
MEDJEE TUTORIALS - Result Video-Analysis deck generator.

Usage:
    python generate.py input.json output.pptx

Input JSON schema (see SKILL.md for the human-facing paste format):
{
  "meta": {
    "institute": "MEDJEE TUTORIALS",
    "test_title": "WEEKLY TEST RESULT",
    "batch": "12th JEE",
    "date": "13 June 2026",
    "marking": {"correct": 4, "incorrect": -1},
    "class_total": 14,          # optional
    "present": 13,              # optional
    "default_questions": 25     # optional; per-subject fallback when 'left' missing
  },
  "students": [
    {
      "name": "Divyansh Sri.",
      "absent": false,          # optional
      "previous_marks": 58,     # optional -> improvement delta
      "subjects": [
        {"name": "Physics",   "correct": 10, "incorrect": 11, "left": 4},
        {"name": "Chemistry", "correct": 8,  "incorrect": 15, "left": 2},
        {"name": "Maths",     "correct": 8,  "incorrect": 6,  "left": 11}
      ]
    }
  ]
}
The generator is tolerant: 'left' or 'questions' may be omitted (falls back to
default_questions); 'attempted' is derived as correct+incorrect.
"""

import sys, json, os, tempfile, hashlib
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
DEFAULT_LOGO = os.path.join(HERE, "assets", "logo.png")
LOGO = DEFAULT_LOGO  # overridden per build() via meta["logo"]

# ---------------------------------------------------------------- palette ----
# DARK THEME
BG       = RGBColor(0x16, 0x13, 0x24)   # slide background (dark purple-black)
PURPLE   = RGBColor(0x5A, 0x3A, 0x93)   # deep fill (table header, badge) w/ white text
PURPLE_D = RGBColor(0x24, 0x17, 0x42)   # deep cover background
ACC      = RGBColor(0xC0, 0xAB, 0xF2)   # light-purple accent TEXT on dark
BLUE     = RGBColor(0x5A, 0x9B, 0xF0)   # brightened for dark
GOLD     = RGBColor(0xF7, 0xC1, 0x2E)
INK      = RGBColor(0xF3, 0xF1, 0xFB)   # primary text (near white)
SLATE    = RGBColor(0xA8, 0xA2, 0xC2)   # muted text
GREEN    = RGBColor(0x36, 0xD3, 0x92)   # brightened for dark
RED      = RGBColor(0xFB, 0x6B, 0x74)   # brightened for dark
GRAY     = RGBColor(0x8A, 0x8F, 0xA6)
PANEL    = RGBColor(0x27, 0x21, 0x3C)   # elevated card / alt row
PANEL_D  = RGBColor(0x31, 0x29, 0x4C)   # grand-total / track
LINE     = RGBColor(0x3B, 0x33, 0x55)   # separators
WHITE    = RGBColor(0xFF, 0xFF, 0xFF)   # text on purple fills
GREEN_BG = RGBColor(0x14, 0x3A, 0x2C)
RED_BG   = RGBColor(0x3E, 0x1E, 0x26)
CARD_BD  = RGBColor(0x40, 0x38, 0x60)   # hairline border on cards
CARD_HI  = RGBColor(0x2E, 0x27, 0x46)   # card gradient top
CARD_LO  = RGBColor(0x22, 0x1C, 0x36)   # card gradient bottom
PURPLE_HI = RGBColor(0x6B, 0x47, 0xAE)  # header gradient light
PURPLE_LO = RGBColor(0x45, 0x2A, 0x78)  # header gradient deep
MEDAL = {1: RGBColor(0xF7, 0xC1, 0x2E), 2: RGBColor(0xC9, 0xCD, 0xD8),
         3: RGBColor(0xCE, 0x8E, 0x5A)}

HEAD_FONT = "Calibri"
BODY_FONT = "Calibri"

EMU_IN = 914400
SW, SH = 13.333, 7.5

# ------------------------------------------------------------- shape helpers -
def _set_radius(shape, inches):
    try:
        target = min(inches, shape.width / EMU_IN / 2, shape.height / EMU_IN / 2)
        frac = max(0.0, min(0.5, target / (shape.width / EMU_IN)))
        shape.adjustments[0] = frac
    except Exception:
        pass

def rect(slide, x, y, w, h, fill=None, line=None, line_w=1.0, radius=0.0,
         shadow=False, rounded=None):
    shp_type = MSO_SHAPE.ROUNDED_RECTANGLE if (rounded if rounded is not None else radius > 0) else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(shp_type, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius > 0:
        _set_radius(s, radius)
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line; s.line.width = Pt(line_w)
    s.shadow.inherit = False
    if shadow:
        _soft_shadow(s)
    return s

def _soft_shadow(shape):
    spPr = shape._element.spPr
    effLst = spPr.find(qn('a:effectLst'))
    if effLst is None:
        effLst = spPr.makeelement(qn('a:effectLst'), {})
        spPr.append(effLst)
    sh = effLst.makeelement(qn('a:outerShdw'),
                            {'blurRad': '90000', 'dist': '38100', 'dir': '5400000', 'rotWithShape': '0'})
    clr = sh.makeelement(qn('a:srgbClr'), {'val': '1E1B2E'})
    alpha = clr.makeelement(qn('a:alpha'), {'val': '18000'})
    clr.append(alpha); sh.append(clr); effLst.append(sh)

def grad(shape, c1, c2, ang_deg=90):
    """Replace a shape's fill with a two-stop linear gradient (c1 -> c2)."""
    spPr = shape._element.spPr
    for tag in ('a:solidFill', 'a:noFill', 'a:gradFill', 'a:blipFill', 'a:pattFill'):
        e = spPr.find(qn(tag))
        if e is not None:
            spPr.remove(e)
    g = spPr.makeelement(qn('a:gradFill'), {'rotWithShape': '1'})
    lst = g.makeelement(qn('a:gsLst'), {})
    for pos, col in ((0, c1), (100000, c2)):
        gs = g.makeelement(qn('a:gs'), {'pos': str(pos)})
        ce = g.makeelement(qn('a:srgbClr'), {'val': str(col)})
        gs.append(ce); lst.append(gs)
    g.append(lst)
    g.append(g.makeelement(qn('a:lin'), {'ang': str(int(ang_deg * 60000)), 'scaled': '1'}))
    ln = spPr.find(qn('a:ln'))
    eff = spPr.find(qn('a:effectLst'))
    anchor = ln if ln is not None else eff
    if anchor is not None:
        anchor.addprevious(g)
    else:
        spPr.append(g)
    return shape

def accent(slide, x, y, w=0.62, color=GOLD, h=0.075):
    """Short rounded accent bar used under section titles."""
    rect(slide, x, y, w, h, fill=color, radius=0.04)

def line_h(slide, x, y, w, color=LINE, weight=1.0):
    s = slide.shapes.add_connector(2, Inches(x), Inches(y), Inches(x + w), Inches(y))
    s.line.color.rgb = color; s.line.width = Pt(weight)
    return s

def text(slide, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         wrap=True, space=1.0):
    """runs: str, or list of (txt,size,color,bold[,font]) tuples, or list of such lists (paragraphs)."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = wrap
    tf.auto_size = MSO_AUTO_SIZE.NONE
    for m in ('left', 'right', 'top', 'bottom'):
        setattr(tf, 'margin_' + m, 0)
    tf.vertical_anchor = anchor
    if isinstance(runs, str):
        runs = [[(runs, 18, INK, False)]]
    elif runs and isinstance(runs[0], tuple):
        runs = [runs]
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.line_spacing = space
        for r in para:
            txt, size, color, bold = r[0], r[1], r[2], r[3]
            font = r[4] if len(r) > 4 else BODY_FONT
            run = p.add_run(); run.text = txt
            run.font.size = Pt(size); run.font.bold = bold
            run.font.color.rgb = color; run.font.name = font
    return tb

def logo(slide, x, y, size):
    slide.shapes.add_picture(LOGO, Inches(x), Inches(y), Inches(size), Inches(size))

_AVATAR_TMP = tempfile.mkdtemp(prefix="medjee_av_")
_AV_COLORS = [PURPLE, BLUE, RGBColor(0x2E, 0x8B, 0x9A), RGBColor(0xB4, 0x7A, 0x0C)]

def _circle_png(src_path):
    """Center-crop a photo to a circle on transparent bg; return temp PNG path or None."""
    try:
        im = Image.open(src_path).convert("RGBA")
    except Exception:
        return None
    w, h = im.size
    s = min(w, h)
    im = im.crop(((w - s) // 2, (h - s) // 2, (w - s) // 2 + s, (h - s) // 2 + s))
    D = 512
    im = im.resize((D, D))
    mask = Image.new("L", (D, D), 0)
    ImageDraw.Draw(mask).ellipse((0, 0, D - 1, D - 1), fill=255)
    out = Image.new("RGBA", (D, D), (0, 0, 0, 0))
    out.paste(im, (0, 0), mask)
    p = os.path.join(_AVATAR_TMP, hashlib.md5(src_path.encode()).hexdigest() + ".png")
    out.save(p)
    return p

def avatar(slide, rec, x, y, d):
    """Circular student photo if provided, else an initials badge. d = diameter (in)."""
    ring = 0.055
    # outer ring
    ov = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x - ring), Inches(y - ring),
                                Inches(d + 2 * ring), Inches(d + 2 * ring))
    ov.fill.solid(); ov.fill.fore_color.rgb = WHITE
    ov.line.color.rgb = GOLD; ov.line.width = Pt(2.2); ov.shadow.inherit = False
    _soft_shadow(ov)
    photo = rec.get("photo")
    cpng = _circle_png(photo) if photo and os.path.exists(str(photo)) else None
    if cpng:
        slide.shapes.add_picture(cpng, Inches(x), Inches(y), Inches(d), Inches(d))
    else:
        name = rec.get("name", "").strip()
        parts = [w for w in name.replace(".", " ").split() if w]
        initials = (parts[0][0] + (parts[1][0] if len(parts) > 1 else "")).upper() if parts else "?"
        col = _AV_COLORS[int(hashlib.md5(name.encode()).hexdigest(), 16) % len(_AV_COLORS)]
        c = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
        c.fill.solid(); c.fill.fore_color.rgb = col
        c.line.fill.background(); c.shadow.inherit = False
        text(slide, x, y - 0.02, d, d, [[(initials, int(d * 26), WHITE, True, HEAD_FONT)]],
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

def new_slide(prs, bg=BG):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    r = rect(s, -0.06, -0.06, SW + 0.12, SH + 0.12, fill=bg)
    r._element.addprevious(r._element)  # keep as background
    return s

# --------------------------------------------------------------- analytics --
def pct(n, d):
    return (100.0 * n / d) if d else 0.0

def fmt_pct(n, d, dash_zero=False):
    if not d:
        return "\u2014" if dash_zero else "0%"
    v = 100.0 * n / d
    return f"{v:.0f}%" if abs(v - round(v)) < 0.05 else f"{v:.1f}%"

def compute(meta, students):
    pc = meta["marking"]["correct"]
    pw = meta["marking"]["incorrect"]
    defq = meta.get("default_questions", 25)
    recs = []
    for st in students:
        rec = {"name": st["name"], "absent": bool(st.get("absent")),
               "previous_marks": st.get("previous_marks"),
               "photo": st.get("photo"), "subjects": []}
        tot = dict(q=0, c=0, w=0, l=0, at=0, m=0, mx=0)
        for sub in st.get("subjects", []):
            c = int(sub.get("correct", 0)); w = int(sub.get("incorrect", 0))
            at = c + w
            if "left" in sub and sub["left"] is not None:
                l = int(sub["left"]); q = at + l
            elif "questions" in sub and sub["questions"] is not None:
                q = int(sub["questions"]); l = max(0, q - at)
            else:
                q = max(defq, at); l = max(0, q - at)
            m = c * pc + w * pw
            mx = q * pc
            s = dict(name=sub["name"], q=q, c=c, w=w, l=l, at=at, m=m, mx=mx,
                     acc=pct(c, at), attpct=pct(at, q), scorepct=pct(m, mx))
            rec["subjects"].append(s)
            for k, v in dict(q=q, c=c, w=w, l=l, at=at, m=m, mx=mx).items():
                tot[k] += v
        rec["tot"] = tot
        rec["accuracy"]  = pct(tot["c"], tot["at"])
        rec["attempt"]   = pct(tot["at"], tot["q"])
        rec["percentage"] = pct(tot["m"], tot["mx"])
        rec["negatives"] = tot["w"] * abs(pw)
        rec["positives"] = tot["c"] * pc
        rec["delta"] = (tot["m"] - rec["previous_marks"]) if rec["previous_marks"] is not None else None
        if rec["subjects"]:
            best = max(rec["subjects"], key=lambda s: s["scorepct"])
            worst = min(rec["subjects"], key=lambda s: s["scorepct"])
            rec["best"], rec["worst"] = best, worst
        else:
            rec["best"] = rec["worst"] = None
        recs.append(rec)
    # ranking over present students
    present = [r for r in recs if not r["absent"]]
    present.sort(key=lambda r: r["tot"]["m"], reverse=True)
    n = len(present)
    for i, r in enumerate(present):
        r["rank"] = i + 1
        r["rank_of"] = n
        atmost = sum(1 for o in present if o["tot"]["m"] <= r["tot"]["m"])
        r["percentile"] = round(100.0 * atmost / n) if n else 0
    if present:
        avg_m = sum(r["tot"]["m"] for r in present) / n
        avg_mx = sum(r["tot"]["mx"] for r in present) / n
        klass = dict(n=n, avg_marks=avg_m, avg_mx=avg_mx,
                     avg_pct=pct(avg_m, avg_mx),
                     topper=present[0]["name"], top_marks=present[0]["tot"]["m"],
                     top_mx=present[0]["tot"]["mx"],
                     avg_acc=sum(r["accuracy"] for r in present) / n)
    else:
        klass = dict(n=0, avg_marks=0, avg_mx=0, avg_pct=0, topper="\u2014",
                     top_marks=0, top_mx=0, avg_acc=0)
    # per-subject class averages (present students only)
    ssum, scnt, sacc, smax = {}, {}, {}, {}
    order = []
    for r in present:
        for sub in r["subjects"]:
            nm = sub["name"]
            if nm not in ssum:
                order.append(nm)
            ssum[nm] = ssum.get(nm, 0) + sub["m"]
            scnt[nm] = scnt.get(nm, 0) + 1
            sacc[nm] = sacc.get(nm, 0) + sub["acc"]
            smax[nm] = sub["mx"]
    klass["subj_order"] = order
    klass["subj_avg"] = {k: ssum[k] / scnt[k] for k in ssum}
    klass["subj_avg_acc"] = {k: sacc[k] / scnt[k] for k in sacc}
    klass["subj_max"] = smax
    imp = [r for r in present if r["delta"] is not None]
    klass["most_improved"] = max(imp, key=lambda r: r["delta"]) if imp else None
    klass["best_acc"] = max(present, key=lambda r: r["accuracy"]) if present else None
    klass["highest"] = present[0] if present else None
    # attach per-subject vs-class delta onto each student's subjects
    for r in present:
        for sub in r["subjects"]:
            sub["cls_avg"] = klass["subj_avg"].get(sub["name"])
            sub["vs_cls"] = (sub["m"] - sub["cls_avg"]) if (len(present) >= 2 and sub["cls_avg"] is not None) else None
    return recs, present, klass

# ------------------------------------------------------------ slide: header --
def header(slide, meta):
    logo(slide, 0.55, 0.42, 0.62)
    text(slide, 1.28, 0.44, 5.0, 0.6,
         [[(meta.get("institute", "MEDJEE TUTORIALS"), 15, ACC, True, HEAD_FONT)],
          [(meta.get("test_title", "TEST RESULT").title() + "  \u2022  " +
            f'{meta.get("batch","")}', 9.5, SLATE, False)]],
         anchor=MSO_ANCHOR.MIDDLE, space=1.0)
    text(slide, SW - 4.05, 0.5, 3.5, 0.45,
         [[(meta.get("date", ""), 12, SLATE, True)]],
         align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    line_h(slide, 0.55, 1.28, SW - 1.10, color=LINE, weight=1.0)

# --------------------------------------------------------------- cover slide -
def slide_cover(prs, meta, klass):
    s = new_slide(prs, bg=PURPLE_D)
    grad(rect(s, -0.06, -0.06, SW + 0.12, SH + 0.12, fill=PURPLE_D),
         RGBColor(0x2E, 0x1D, 0x52), RGBColor(0x18, 0x0F, 0x2E), ang_deg=115)
    # subtle geometric accents
    grad(rect(s, 9.4, -1.6, 6.0, 6.0, fill=PURPLE, radius=3.0),
         RGBColor(0x4C, 0x2B, 0x82), RGBColor(0x2C, 0x1A, 0x54), ang_deg=90)
    rect(s, -1.6, 4.7, 5.2, 4.2, fill=RGBColor(0x3A, 0x22, 0x66), radius=2.4)
    logo(s, 0.7, 0.6, 1.4)
    text(s, 2.3, 0.7, 8.0, 1.25,
         [[(meta.get("institute", "MEDJEE TUTORIALS"), 22, WHITE, True, HEAD_FONT)],
          [("RESULT VIDEO ANALYSIS", 12.5, GOLD, True)]],
         anchor=MSO_ANCHOR.MIDDLE, space=1.1)
    accent(s, 0.72, 2.5, 0.9, GOLD, 0.09)
    text(s, 0.7, 2.66, 11.9, 1.5,
         [[(meta.get("test_title", "TEST RESULT").upper(), 46, WHITE, True, HEAD_FONT)]],
         anchor=MSO_ANCHOR.MIDDLE)
    _sub = meta.get("batch", "")
    if meta.get("date"):
        _sub += "    \u2022    " + meta["date"]
    text(s, 0.72, 3.98, 11.9, 0.7,
         [[(_sub, 18, RGBColor(0xCD, 0xC3, 0xE6), True)]])
    # stat chips
    chips = []
    if meta.get("class_total") is not None:
        chips.append(("STUDENTS", str(meta.get("class_total"))))
    if meta.get("present") is not None:
        chips.append(("PRESENT", str(meta.get("present"))))
    if klass["n"]:
        chips.append(("CLASS AVERAGE", f'{klass["avg_pct"]:.1f}%'))
        chips.append(("TOPPER", klass["topper"]))
    cw, gap, y = 2.72, 0.30, 5.2
    total = len(chips) * cw + (len(chips) - 1) * gap
    x = (SW - total) / 2
    for label, val in chips:
        card = rect(s, x, y, cw, 1.35, fill=PURPLE, radius=0.14,
                    line=RGBColor(0x59, 0x3A, 0x8E), line_w=1.0, shadow=True)
        grad(card, RGBColor(0x4A, 0x2C, 0x80), RGBColor(0x33, 0x1E, 0x5C), ang_deg=90)
        rect(s, x, y, 0.12, 1.35, fill=GOLD, radius=0.06)
        text(s, x + 0.30, y + 0.22, cw - 0.42, 0.4, [[(label, 10, GOLD, True)]])
        vsize = 24 if len(val) <= 8 else (16 if len(val) <= 14 else 13)
        text(s, x + 0.28, y + 0.55, cw - 0.44, 0.7, [[(val, vsize, WHITE, True, HEAD_FONT)]],
             anchor=MSO_ANCHOR.MIDDLE)
        x += cw + gap
    text(s, 0.7, SH - 0.62, 11.9, 0.4,
         [[("Prepared by MEDJEE TUTORIALS  \u2022  Keep pushing forward", 10.5,
            RGBColor(0x9C, 0x90, 0xBE), False)]])
    return s

# ---------------------------------------------------- slide: student scorecard
def slide_scorecard(prs, meta, rec):
    s = new_slide(prs)
    header(s, meta)
    if rec["absent"]:
        return _absent_card(s, rec)
    top = 1.62
    # avatar + name + rank
    av_d = 0.95
    avatar(s, rec, 0.6, top + (0.85 - av_d) / 2, av_d)
    name_x = 0.6 + av_d + 0.32
    text(s, name_x, top - 0.02, SW - 0.55 - 2.75 - name_x, 0.72,
         [[(rec["name"], 30, INK, True, HEAD_FONT)]], anchor=MSO_ANCHOR.MIDDLE)
    accent(s, name_x + 0.02, top + 0.70, 0.75, GOLD, 0.07)
    rk = f'RANK  {rec["rank"]} / {rec["rank_of"]}'
    rw = 2.55
    badge = rect(s, SW - 0.55 - rw, top + 0.06, rw, 0.74, fill=PURPLE, radius=0.15,
                 line=RGBColor(0x6A, 0x49, 0xA8), line_w=1.0, shadow=True)
    grad(badge, PURPLE_HI, PURPLE_LO, ang_deg=90)
    if rec["rank"] in MEDAL:
        rect(s, SW - 0.55 - rw + 0.16, top + 0.06 + 0.20, 0.34, 0.34,
             fill=MEDAL[rec["rank"]], radius=0.17)
        text(s, SW - 0.55 - rw + 0.16, top + 0.045 + 0.20, 0.34, 0.34,
             [[(str(rec["rank"]), 14, RGBColor(0x2A, 0x20, 0x12), True, HEAD_FONT)]],
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        text(s, SW - 0.55 - rw + 0.42, top + 0.06, rw - 0.42, 0.74,
             [[(rk, 16, WHITE, True, HEAD_FONT)]], align=PP_ALIGN.CENTER,
             anchor=MSO_ANCHOR.MIDDLE)
    else:
        text(s, SW - 0.55 - rw, top + 0.06, rw, 0.74, [[(rk, 16, WHITE, True, HEAD_FONT)]],
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    body_y = top + 1.05
    # ---- left hero score card ----
    lx, lw = 0.55, 4.15
    lh = SH - body_y - 0.55
    hero = rect(s, lx, body_y, lw, lh, fill=PANEL, radius=0.16,
                line=CARD_BD, line_w=1.0, shadow=True)
    grad(hero, CARD_HI, CARD_LO, ang_deg=90)
    rect(s, lx, body_y, lw, 0.14, fill=GOLD, radius=0.06)
    tot = rec["tot"]
    text(s, lx, body_y + 0.32, lw, 0.4, [[("TOTAL SCORE", 13, SLATE, True)]],
         align=PP_ALIGN.CENTER)
    text(s, lx, body_y + 0.60, lw, 1.25,
         [[(str(tot["m"]), 78, ACC, True, HEAD_FONT),
           (f' / {tot["mx"]}', 28, SLATE, True, HEAD_FONT)]],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    # percentage bar
    bx, bw, by = lx + 0.45, lw - 0.9, body_y + 2.05
    text(s, lx, by - 0.38, lw, 0.34,
         [[(f'{rec["percentage"]:.1f}% SCORE', 15, INK, True)]], align=PP_ALIGN.CENTER)
    rect(s, bx, by, bw, 0.28, fill=PANEL_D, radius=0.14)
    fillw = max(0.0, min(1.0, rec["percentage"] / 100.0)) * bw
    if fillw > 0.05:
        pbar = rect(s, bx, by, fillw, 0.28, fill=GREEN, radius=0.14)
        if rec["percentage"] >= 33:
            grad(pbar, RGBColor(0x3C, 0xE0, 0x9A), RGBColor(0x18, 0xA6, 0x70), ang_deg=0)
        else:
            grad(pbar, RGBColor(0xFB, 0x7B, 0x84), RGBColor(0xD8, 0x38, 0x44), ang_deg=0)
    # mini stats
    stat_pairs = [("Accuracy", f'{rec["accuracy"]:.0f}%'),
                  ("Attempt", f'{rec["attempt"]:.0f}%'),
                  ("Percentile", f'{rec["percentile"]}')]
    if rec["delta"] is not None:
        d = rec["delta"]
        stat_pairs[2] = ("Improvement", (f'+{d}' if d >= 0 else f'{d}'))
    sy = by + 0.60
    cellw = (lw - 0.5) / 3
    for i, (lab, val) in enumerate(stat_pairs):
        cx = lx + 0.25 + i * cellw
        vcol = INK
        if lab == "Improvement":
            vcol = GREEN if rec["delta"] >= 0 else RED
        text(s, cx, sy, cellw, 0.5, [[(val, 25, vcol, True, HEAD_FONT)]],
             align=PP_ALIGN.CENTER)
        text(s, cx, sy + 0.50, cellw, 0.3, [[(lab.upper(), 9.5, SLATE, True)]],
             align=PP_ALIGN.CENTER)
    # verdict strip
    if rec["best"] and rec["worst"]:
        vy = body_y + lh - 0.84
        line_h(s, lx + 0.3, vy - 0.14, lw - 0.6, color=LINE)
        rect(s, lx + 0.3, vy + 0.04, 0.16, 0.16, fill=GREEN, radius=0.03)
        text(s, lx + 0.56, vy, lw - 0.8, 0.32,
             [[("Strong  ", 11, SLATE, True), (rec["best"]["name"], 11, GREEN, True)]],
             anchor=MSO_ANCHOR.MIDDLE)
        rect(s, lx + 0.3, vy + 0.40, 0.16, 0.16, fill=RED, radius=0.03)
        text(s, lx + 0.56, vy + 0.36, lw - 0.8, 0.32,
             [[("Focus   ", 11, SLATE, True), (rec["worst"]["name"], 11, RED, True)]],
             anchor=MSO_ANCHOR.MIDDLE)

    # ---- right subject table ----
    tx = lx + lw + 0.4
    tw = SW - 0.55 - tx
    _subject_table(s, rec, tx, body_y, tw, lh)
    return s

def _subject_table(s, rec, x, y, w, h):
    cols = ["SUBJECT", "Q", "ATT", "CORR", "INCORR", "LEFT", "ACC", "SCORE"]
    weights = [1.55, 0.55, 0.62, 0.7, 0.78, 0.62, 0.72, 1.0]
    tw = sum(weights)
    widths = [wt / tw * w for wt in weights]
    xs = [x]
    for wd in widths:
        xs.append(xs[-1] + wd)
    n = len(rec["subjects"])
    head_h = 0.54
    row_h = min(0.72, (h - head_h - 0.72) / max(1, n))
    # header row
    hdr = rect(s, x, y, w, head_h, fill=PURPLE, radius=0.10)
    grad(hdr, PURPLE_HI, PURPLE_LO, ang_deg=90)
    for i, c in enumerate(cols):
        al = PP_ALIGN.LEFT if i == 0 else PP_ALIGN.CENTER
        pad = 0.14 if i == 0 else 0
        text(s, xs[i] + pad, y, widths[i] - pad, head_h, [[(c, 11, WHITE, True)]],
             align=al, anchor=MSO_ANCHOR.MIDDLE)
    ry = y + head_h + 0.06
    for k, sub in enumerate(rec["subjects"]):
        if k % 2 == 1:
            rect(s, x, ry, w, row_h, fill=PANEL)
        vals = [sub["name"], str(sub["q"]), str(sub["at"]), str(sub["c"]),
                str(sub["w"]), str(sub["l"]), f'{sub["acc"]:.0f}%', f'{sub["m"]}/{sub["mx"]}']
        for i, v in enumerate(vals):
            al = PP_ALIGN.LEFT if i == 0 else PP_ALIGN.CENTER
            pad = 0.14 if i == 0 else 0
            col = INK; bold = (i == 0 or i == 7)
            if i == 3: col = GREEN
            if i == 4: col = RED
            if i == 6: col = GREEN if sub["acc"] >= 50 else (RED if sub["acc"] < 33 else INK)
            sz = 15 if i == 0 else (17 if i == 7 else 15)
            text(s, xs[i] + pad, ry, widths[i] - pad, row_h, [[(v, sz, col, bold)]],
                 align=al, anchor=MSO_ANCHOR.MIDDLE)
        ry += row_h
    # grand total row
    line_h(s, x, ry + 0.02, w, color=SLATE, weight=1.3)
    ry += 0.12
    tot = rec["tot"]
    gt = ["GRAND TOTAL", str(tot["q"]), str(tot["at"]), str(tot["c"]),
          str(tot["w"]), str(tot["l"]), f'{rec["accuracy"]:.0f}%',
          f'{tot["m"]}/{tot["mx"]}']
    rect(s, x, ry, w, 0.66, fill=PANEL_D, radius=0.10)
    for i, v in enumerate(gt):
        al = PP_ALIGN.LEFT if i == 0 else PP_ALIGN.CENTER
        pad = 0.14 if i == 0 else 0
        col = ACC if i in (0, 7) else INK
        sz = 15 if i == 0 else (17 if i == 7 else 15)
        text(s, xs[i] + pad, ry, widths[i] - pad, 0.66, [[(v, sz, col, True)]],
             align=al, anchor=MSO_ANCHOR.MIDDLE)

def _absent_card(s, rec):
    text(s, 0.55, 1.62, 10, 0.85, [[(rec["name"], 30, INK, True, HEAD_FONT)]],
         anchor=MSO_ANCHOR.MIDDLE)
    cy = 3.0
    rect(s, 3.9, cy, 5.5, 2.2, fill=RED_BG, radius=0.18, shadow=True)
    text(s, 3.9, cy + 0.5, 5.5, 0.7, [[("ABSENT", 40, RED, True, HEAD_FONT)]],
         align=PP_ALIGN.CENTER)
    text(s, 3.9, cy + 1.4, 5.5, 0.5,
         [[("Not present for this test", 14, SLATE, True)]], align=PP_ALIGN.CENTER)
    return s

# ------------------------------------------------ slide: performance analytics
def slide_analytics(prs, meta, rec, klass=None):
    if rec["absent"]:
        return None
    s = new_slide(prs)
    header(s, meta)
    text(s, 0.55, 1.5, 9, 0.62, [[("PERFORMANCE ANALYTICS", 24, INK, True, HEAD_FONT)]],
         anchor=MSO_ANCHOR.MIDDLE)
    accent(s, 0.57, 2.14, 0.72, GOLD, 0.07)
    text(s, SW - 4.55, 1.55, 4.0, 0.6, [[(rec["name"], 16, ACC, True, HEAD_FONT)]],
         align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    # insight line
    tot = rec["tot"]
    if rec["best"] and rec["worst"]:
        text(s, SW - 8.55, 1.95, 8.0, 0.34,
             [[("Strongest ", 10.5, SLATE, True), (rec["best"]["name"], 10.5, GREEN, True),
               ("   \u2022   Focus ", 10.5, SLATE, True), (rec["worst"]["name"], 10.5, RED, True),
               (f"   \u2022   +{rec['positives']} earned, ", 10.5, SLATE, True),
               (f"-{rec['negatives']} lost", 10.5, RED, True)]],
             align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)

    # ---- KPI cards ----
    kpis = [("OVERALL ACCURACY", f'{rec["accuracy"]:.0f}%',
             "of attempted correct", GREEN if rec["accuracy"] >= 50 else RED),
            ("ATTEMPT RATE", f'{rec["attempt"]:.0f}%',
             f'{tot["at"]} of {tot["q"]} qs', BLUE),
            ("NEGATIVE MARKS", f'-{rec["negatives"]}',
             f'{tot["w"]} wrong attempts', RED),
            ("PERCENTILE", f'{rec["percentile"]}',
             f'rank {rec["rank"]} of {rec["rank_of"]}', ACC)]
    ky = 2.5; kh = 1.55
    gap = 0.28
    kw = (SW - 1.1 - 3 * gap) / 4
    x = 0.55
    for lab, val, sub, col in kpis:
        card = rect(s, x, ky, kw, kh, fill=PANEL, radius=0.14,
                    line=CARD_BD, line_w=1.0, shadow=True)
        grad(card, CARD_HI, CARD_LO, ang_deg=90)
        rect(s, x, ky, 0.13, kh, fill=col, radius=0.06)
        text(s, x + 0.30, ky + 0.22, kw - 0.42, 0.35, [[(lab, 10, SLATE, True)]])
        text(s, x + 0.28, ky + 0.48, kw - 0.44, 0.66, [[(val, 38, col, True, HEAD_FONT)]],
             anchor=MSO_ANCHOR.MIDDLE)
        text(s, x + 0.30, ky + 1.18, kw - 0.42, 0.3, [[(sub, 9.5, SLATE, False)]])
        x += kw + gap

    # ---- subject composition bars ----
    by = ky + kh + 0.5
    text(s, 0.55, by, 8, 0.4, [[("SUBJECT-WISE ATTEMPT COMPOSITION", 13, INK, True)]])
    # legend
    _legend(s, SW - 5.0, by + 0.02, [("Correct", GREEN), ("Incorrect", RED), ("Left", GRAY)])
    by += 0.55
    rows = rec["subjects"]
    avail = SH - by - 0.55
    rh = min(0.86, avail / max(1, len(rows)))
    label_w = 1.7
    barx = 0.55 + label_w + 0.15
    barw = SW - 0.55 - barx - 2.5  # room for acc + score/vs-avg at right
    for sub in rows:
        cy = by + rh * 0.16
        bh = rh * 0.5
        text(s, 0.55, by, label_w, rh, [[(sub["name"], 14, INK, True)]],
             anchor=MSO_ANCHOR.MIDDLE)
        rect(s, barx, cy, barw, bh, fill=PANEL, radius=0.05)
        q = max(1, sub["q"])
        segs = [(sub["c"], GREEN), (sub["w"], RED), (sub["l"], GRAY)]
        cx = barx
        for cnt, col in segs:
            sw = cnt / q * barw
            if sw > 0.02:
                rect(s, cx, cy, sw, bh, fill=col)
                if sw > 0.32:
                    text(s, cx, cy, sw, bh, [[(str(cnt), 13, RGBColor(0x12,0x10,0x1E), True)]],
                         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            cx += sw
        # accuracy + score + vs-class at far right
        text(s, SW - 0.55 - 2.42, by, 0.86, rh,
             [[(f'{sub["acc"]:.0f}%', 16, GREEN if sub["acc"] >= 50 else RED, True, HEAD_FONT)],
              [("acc", 8.5, SLATE, True)]], align=PP_ALIGN.CENTER,
             anchor=MSO_ANCHOR.MIDDLE, space=0.95)
        score_para = [[(f'{sub["m"]}', 17, INK, True, HEAD_FONT), (f'/{sub["mx"]}', 11, SLATE, True)]]
        if sub.get("vs_cls") is not None:
            d = round(sub["vs_cls"])
            arrow = "\u25B2" if d > 0 else ("\u25BC" if d < 0 else "\u2013")
            dcol = GREEN if d > 0 else (RED if d < 0 else SLATE)
            dtxt = f'{arrow} {"+" if d > 0 else ""}{d} vs avg'
            score_para.append([(dtxt, 9.5, dcol, True)])
        text(s, SW - 0.55 - 1.52, by, 1.52, rh, score_para,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, space=0.95)
        by += rh
    return s

def _legend(s, x, y, items):
    for lab, col in items:
        rect(s, x, y + 0.07, 0.2, 0.2, fill=col, radius=0.03)
        text(s, x + 0.28, y, 1.1, 0.34, [[(lab, 10.5, SLATE, True)]],
             anchor=MSO_ANCHOR.MIDDLE)
        x += 1.15 + 0.12 * len(lab) / 3

# ----------------------------------------------------- slide: class leaderboard
def slide_leaderboard(prs, meta, present, klass):
    if not present:
        return None
    s = new_slide(prs)
    header(s, meta)
    text(s, 0.55, 1.5, 9, 0.62, [[("CLASS LEADERBOARD", 24, INK, True, HEAD_FONT)]],
         anchor=MSO_ANCHOR.MIDDLE)
    accent(s, 0.57, 2.14, 0.72, GOLD, 0.07)
    text(s, SW - 5.05, 1.55, 4.5, 0.6,
         [[("Class Average  ", 12, SLATE, True),
           (f'{klass["avg_pct"]:.1f}%', 18, ACC, True, HEAD_FONT)]],
         align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)

    x, y = 0.55, 2.45
    w = SW - 1.1
    cols = ["RANK", "STUDENT", "SCORE", "PERCENTAGE", "ACCURACY", "ATTEMPT"]
    weights = [0.9, 3.4, 1.4, 1.5, 1.3, 1.3]
    tw = sum(weights); widths = [wt / tw * w for wt in weights]
    xs = [x]
    for wd in widths:
        xs.append(xs[-1] + wd)
    head_h = 0.58
    grad(rect(s, x, y, w, head_h, fill=PURPLE, radius=0.10), PURPLE_HI, PURPLE_LO, ang_deg=90)
    for i, c in enumerate(cols):
        al = PP_ALIGN.LEFT if i == 1 else PP_ALIGN.CENTER
        pad = 0.18 if i == 1 else 0
        text(s, xs[i] + pad, y, widths[i] - pad, head_h, [[(c, 12, WHITE, True)]],
             align=al, anchor=MSO_ANCHOR.MIDDLE)
    ry = y + head_h + 0.08
    n = len(present)
    avail = SH - ry - 0.55
    rh = min(0.74, avail / max(1, n))
    tint = {1: RGBColor(0x40,0x35,0x16), 2: RGBColor(0x31,0x33,0x3E), 3: RGBColor(0x3E,0x2F,0x22)}
    for r in present:
        rk = r["rank"]
        if rk in tint:
            row = rect(s, x, ry, w, rh, fill=tint[rk], radius=0.06)
            grad(row, tint[rk], BG, ang_deg=0)
        elif rk % 2 == 0:
            rect(s, x, ry, w, rh, fill=PANEL)
        vals = [None, r["name"], f'{r["tot"]["m"]}/{r["tot"]["mx"]}',
                f'{r["percentage"]:.1f}%', f'{r["accuracy"]:.0f}%', f'{r["attempt"]:.0f}%']
        # rank cell: medal disc for top-3, else number
        if rk in MEDAL:
            md = min(0.44, rh - 0.16)
            cxm = xs[0] + widths[0] / 2 - md / 2
            rect(s, cxm, ry + (rh - md) / 2, md, md, fill=MEDAL[rk], radius=md / 2)
            text(s, cxm, ry + (rh - md) / 2 - 0.01, md, md,
                 [[(str(rk), 16, RGBColor(0x2A, 0x20, 0x12), True, HEAD_FONT)]],
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        else:
            text(s, xs[0], ry, widths[0], rh, [[(str(rk), 15, INK, True)]],
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        for i, v in enumerate(vals):
            if v is None:
                continue
            al = PP_ALIGN.LEFT if i == 1 else PP_ALIGN.CENTER
            pad = 0.18 if i == 1 else 0
            bold = (i in (1, 2))
            col = INK
            if i == 2: col = ACC
            if i == 4: col = GREEN if r["accuracy"] >= 50 else (RED if r["accuracy"] < 33 else INK)
            sz = 16 if i == 1 else (17 if i == 2 else 15)
            text(s, xs[i] + pad, ry, widths[i] - pad, rh, [[(v, sz, col, bold)]],
                 align=al, anchor=MSO_ANCHOR.MIDDLE)
        ry += rh
    return s

# ------------------------------------------------------ slide: class insights
def slide_class_insights(prs, meta, present, klass):
    if len(present) < 2:
        return None
    s = new_slide(prs)
    header(s, meta)
    text(s, 0.55, 1.5, 9, 0.62, [[("CLASS INSIGHTS", 24, INK, True, HEAD_FONT)]],
         anchor=MSO_ANCHOR.MIDDLE)
    accent(s, 0.57, 2.14, 0.72, GOLD, 0.07)
    text(s, SW - 5.05, 1.55, 4.5, 0.6,
         [[("Class Average  ", 12, SLATE, True),
           (f'{klass["avg_pct"]:.1f}%', 18, ACC, True, HEAD_FONT)]],
         align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)

    # ---- left: podium (top 3) ----
    top = present[:3]
    px, pw = 0.55, 4.2
    py, ph = 2.55, 4.35
    panel = rect(s, px, py, pw, ph, fill=PANEL, radius=0.16, line=CARD_BD, line_w=1.0, shadow=True)
    grad(panel, CARD_HI, CARD_LO, 90)
    text(s, px, py + 0.24, pw, 0.4, [[("TOP PERFORMERS", 12, SLATE, True)]],
         align=PP_ALIGN.CENTER)
    rowh = (ph - 0.8) / max(1, len(top))
    yy = py + 0.72
    for r in top:
        md = 0.52
        rect(s, px + 0.32, yy + rowh / 2 - md / 2, md, md, fill=MEDAL[r["rank"]], radius=md / 2)
        text(s, px + 0.32, yy + rowh / 2 - md / 2 - 0.01, md, md,
             [[(str(r["rank"]), 19, RGBColor(0x2A, 0x20, 0x12), True, HEAD_FONT)]],
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        text(s, px + 1.05, yy, pw - 2.1, rowh,
             [[(r["name"], 16, INK, True, HEAD_FONT)],
              [(f'{r["percentage"]:.1f}%  \u2022  {r["accuracy"]:.0f}% acc', 10.5, SLATE, True)]],
             anchor=MSO_ANCHOR.MIDDLE, space=1.0)
        text(s, px + pw - 1.35, yy, 1.15, rowh,
             [[(str(r["tot"]["m"]), 22, ACC, True, HEAD_FONT)],
              [(f'/{r["tot"]["mx"]}', 10, SLATE, True)]],
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, space=0.9)
        if r["rank"] != top[-1]["rank"]:
            line_h(s, px + 0.3, yy + rowh, pw - 0.6, color=LINE)
        yy += rowh

    # ---- right: subject-wise class average bars ----
    rx = px + pw + 0.4
    rw = SW - 0.55 - rx
    text(s, rx, py, rw, 0.4, [[("SUBJECT-WISE CLASS AVERAGE", 13, INK, True)]])
    order = klass.get("subj_order", [])
    n_sub = max(1, len(order))
    zone_y = py + 0.6
    zone_h = ph - 0.6
    srh = min(1.15, zone_h / n_sub)
    barx = rx + 1.7
    barw = rw - 1.7 - 1.35
    for nm in order:
        avg = klass["subj_avg"].get(nm, 0)
        mx = klass["subj_max"].get(nm, 1) or 1
        aacc = klass["subj_avg_acc"].get(nm, 0)
        frac = max(0.0, min(1.0, avg / mx))
        cyc = zone_y + srh / 2
        text(s, rx, zone_y, 1.65, srh, [[(nm, 14, INK, True)]], anchor=MSO_ANCHOR.MIDDLE)
        rect(s, barx, cyc - 0.14, barw, 0.28, fill=PANEL_D, radius=0.14)
        fillw = frac * barw
        bcol = GREEN if frac >= 0.5 else (GOLD if frac >= 0.33 else RED)
        if fillw > 0.05:
            grad(rect(s, barx, cyc - 0.14, fillw, 0.28, fill=bcol, radius=0.14),
                 bcol, RGBColor(max(bcol[0]-40,0), max(bcol[1]-40,0), max(bcol[2]-40,0)), 0)
        text(s, barx + barw + 0.12, zone_y, 1.2, srh,
             [[(f'{avg:.0f}', 17, INK, True, HEAD_FONT), (f'/{mx}', 10, SLATE, True)],
              [(f'{aacc:.0f}% acc', 9.5, SLATE, True)]],
             anchor=MSO_ANCHOR.MIDDLE, space=0.92)
        zone_y += srh
    return s

# --------------------------------------------------------------- pdf + split -
def _to_pdf(pptx_path, pdf_path):
    import subprocess, shutil, tempfile
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if not exe:
        raise RuntimeError("LibreOffice (soffice) not found for PDF export")
    td = tempfile.mkdtemp()
    subprocess.run([exe, "--headless", "--convert-to", "pdf", "--outdir", td, pptx_path],
                   check=True, timeout=240,
                   stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    src = os.path.join(td, os.path.splitext(os.path.basename(pptx_path))[0] + ".pdf")
    shutil.move(src, pdf_path)

def _is_neet(students):
    keys = ("bio", "botan", "zoolog")
    for st in students:
        for sub in st.get("subjects", []):
            if any(k in sub["name"].lower() for k in keys):
                return True
    return False

def _auto_batch(meta, students):
    base = meta.get("batch", "").strip()
    return base if base else ("NEET" if _is_neet(students) else "JEE")

def build_batches(data, outdir, as_pdf=False):
    """Split students into batches (by their 'batch' field, else auto JEE/NEET) and
    emit one deck per batch. Returns list of output paths."""
    meta0 = data["meta"]
    groups = {}
    for st in data["students"]:
        b = st.get("batch") or _auto_batch(meta0, [st])
        groups.setdefault(b, []).append(st)
    os.makedirs(outdir, exist_ok=True)
    outs = []
    for batch, sts in groups.items():
        meta = dict(meta0)
        meta["batch"] = batch
        meta["default_questions"] = 45 if _is_neet(sts) else 25
        meta["class_total"] = len(sts)
        meta["present"] = len([x for x in sts if not x.get("absent")])
        title = meta.get("test_title", "Weekly Test Result")
        date = meta.get("date", "")
        stem = f'{batch} - {title} {date} - Analysis'.strip().replace("  ", " ")
        pptx = os.path.join(outdir, stem + ".pptx")
        build({"meta": meta, "students": sts}, pptx)
        if as_pdf:
            pdf = os.path.join(outdir, stem + ".pdf")
            _to_pdf(pptx, pdf); os.remove(pptx); outs.append(pdf)
        else:
            outs.append(pptx)
    return outs

# --------------------------------------------------------------------- build --
def build(data, out_path):
    meta = data["meta"]
    meta.setdefault("institute", "MEDJEE TUTORIALS")
    meta.setdefault("marking", {"correct": 4, "incorrect": -1})
    global LOGO
    LOGO = meta.get("logo") or DEFAULT_LOGO
    students = data["students"]
    recs, present, klass = compute(meta, students)

    prs = Presentation()
    prs.slide_width = Inches(SW); prs.slide_height = Inches(SH)

    slide_cover(prs, meta, klass)
    slide_class_insights(prs, meta, present, klass)
    # students in rank order (present first, then absentees at end)
    ordered = sorted(recs, key=lambda r: (r["absent"], r.get("rank", 9999)))
    for rec in ordered:
        slide_scorecard(prs, meta, rec)
        slide_analytics(prs, meta, rec, klass)
    slide_leaderboard(prs, meta, present, klass)

    # if a .pdf path is requested, render pptx to temp then convert
    if out_path.lower().endswith(".pdf"):
        import tempfile
        tmp = os.path.join(tempfile.mkdtemp(), "deck.pptx")
        prs.save(tmp); _to_pdf(tmp, out_path)
    else:
        prs.save(out_path)
    return len(prs.slides._sldIdLst)

def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    flags = {a for a in sys.argv[1:] if a.startswith("--")}
    if len(args) < 2:
        print("usage: python generate.py input.json output.pptx|output.pdf|OUTDIR [--split] [--pdf]")
        sys.exit(1)
    with open(args[0], encoding="utf-8") as f:
        data = json.load(f)
    if "--split" in flags:
        outs = build_batches(data, args[1], as_pdf=("--pdf" in flags))
        print("OK  " + str(len(outs)) + " deck(s):")
        for o in outs:
            print("  " + o)
    else:
        out = args[1]
        if "--pdf" in flags and not out.lower().endswith(".pdf"):
            out = os.path.splitext(out)[0] + ".pdf"
        n = build(data, out)
        print(f"OK  {out}  ({n} slides)")

if __name__ == "__main__":
    main()
